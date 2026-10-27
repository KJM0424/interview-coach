"""AI 면접 코치 Streamlit 앱."""

from __future__ import annotations

import io
import json
import os
import re
from collections.abc import Callable
from html import escape
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from dotenv import load_dotenv
from pypdf import PdfReader

import llm_client
import prompts
from llm_client import LLMClient, LLMError, LLMResponseError

PASS_SCORE = 6.0
MAX_REANSWERS = 2
ROLE_OPTIONS = ["AI 엔지니어", "백엔드 개발자", "직접 입력"]
CUSTOM_ROLE = "직접 입력"
SUPPORTED_TYPES = ["txt", "md", "pdf", "docx"]


# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------


def read_setting(name: str) -> str | None:
    """st.secrets를 먼저 읽고, 없으면 .env(환경변수)에서 읽는다."""
    try:
        value = st.secrets[name]
    except (KeyError, FileNotFoundError):
        value = None
    if value:
        return str(value)
    load_dotenv()
    return os.getenv(name) or None


def load_provider() -> str:
    """LLM_PROVIDER 설정. gemini(기본) 또는 anthropic."""
    return (read_setting("LLM_PROVIDER") or llm_client.DEFAULT_PROVIDER).strip().lower()


@st.cache_resource
def get_client(provider: str, api_key: str) -> LLMClient:
    return llm_client.create_client(provider, api_key)


# ---------------------------------------------------------------------------
# 업로드 파일 처리
# ---------------------------------------------------------------------------


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text: str) -> str:
    """제어 문자를 공백으로 바꾼다. 브라우저에서 만든 PDF는 띄어쓰기가 \\x01로 추출되기도 한다."""
    return CONTROL_CHARS.sub(" ", text)


def extract_text(filename: str, data: bytes) -> str:
    """업로드한 파일에서 텍스트를 뽑는다."""
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix in ("txt", "md"):
        return clean_text(decode_text(data))
    if suffix == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))
    if suffix == "docx":
        document = Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return clean_text("\n".join(lines))
    raise ValueError(f"지원하지 않는 파일 형식입니다: {filename}")


def read_uploaded_files(files: list[Any]) -> tuple[list[dict[str, str]], list[str]]:
    """업로드 파일 목록을 (자료 목록, 경고 메시지 목록)으로 바꾼다."""
    materials: list[dict[str, str]] = []
    warnings: list[str] = []
    for file in files:
        try:
            text = extract_text(file.name, file.getvalue()).strip()
        except Exception as error:  # 손상된 PDF/DOCX는 라이브러리마다 예외가 달라 한 번에 받는다
            warnings.append(f"{file.name}: 파일을 읽지 못했어요 ({error})")
            continue
        if not text:
            if file.name.lower().endswith(".pdf"):
                warnings.append(f"{file.name}: 글자를 찾지 못했어요. 스캔한 이미지 PDF일 수 있어요")
            else:
                warnings.append(f"{file.name}: 내용이 비어 있어요")
            continue
        materials.append({"name": file.name, "text": text})
    return materials, warnings


# ---------------------------------------------------------------------------
# LLM 호출 (API별 차이는 llm_client.py에서 처리)
# ---------------------------------------------------------------------------


def call_llm_text(client: LLMClient, system: str, user_prompt: str) -> str:
    try:
        return client.complete(system, user_prompt)
    except llm_client.TruncatedResponseError as error:
        raise LLMResponseError(str(error)) from error


def parse_json_object(text: str) -> dict[str, Any]:
    """응답에서 첫 '{'부터 마지막 '}'까지를 JSON 객체로 파싱한다."""
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("응답에서 JSON 객체를 찾지 못했습니다.")
    data = json.loads(match.group(0))
    if not isinstance(data, dict):
        raise ValueError("JSON 최상위 값이 객체가 아닙니다.")
    return data


def call_llm_json(
    client: LLMClient,
    system: str,
    user_prompt: str,
    validate: Callable[[dict[str, Any]], dict[str, Any]],
) -> dict[str, Any]:
    """JSON 응답을 받아 검증한다. 파싱·검증에 실패하면 1번 재시도한다."""
    last_error: Exception | None = None
    for _ in range(2):
        try:
            raw = client.complete(system, user_prompt, json_mode=True)
            return validate(parse_json_object(raw))
        except (ValueError, KeyError, TypeError) as error:
            last_error = error
    raise LLMResponseError(f"JSON 응답 처리에 2번 실패했어요: {last_error}")


def clean_question(text: str) -> str:
    return text.strip().strip('"“”').strip()


# ---------------------------------------------------------------------------
# 응답 검증
# ---------------------------------------------------------------------------


def validate_evaluation(data: dict[str, Any]) -> dict[str, Any]:
    """평가위원 4명의 점수·이유·개선 포인트가 모두 있는지 확인하고 평균을 계산한다."""
    items = data["evaluations"]
    if not isinstance(items, list):
        raise TypeError("evaluations가 리스트가 아닙니다.")
    by_name = {str(item["evaluator"]).strip(): item for item in items}
    evaluations = []
    for evaluator in prompts.EVALUATORS:
        item = by_name[evaluator["name"]]
        score = float(item["score"])
        if not 0 <= score <= 10:
            raise ValueError(f"{evaluator['name']} 점수가 0~10 범위를 벗어났습니다: {score}")
        reason = str(item["reason"]).strip()
        improvement = str(item["improvement"]).strip()
        if not reason or not improvement:
            raise ValueError(f"{evaluator['name']}의 이유 또는 개선 포인트가 비었습니다.")
        evaluations.append(
            {
                "evaluator": evaluator["name"],
                "score": score,
                "reason": reason,
                "improvement": improvement,
            }
        )
    average = sum(e["score"] for e in evaluations) / len(evaluations)
    return {
        "evaluations": evaluations,
        "average": round(average, 1),
        "gap": str(data.get("gap", "")).strip(),
    }


def validate_report(data: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(data["strengths"], list) or not isinstance(data["weaknesses"], list):
        raise TypeError("strengths와 weaknesses는 리스트여야 합니다.")
    strengths = [str(s).strip() for s in data["strengths"] if str(s).strip()]
    weaknesses = [str(w).strip() for w in data["weaknesses"] if str(w).strip()]
    model_answer = str(data["model_answer"]).strip()
    if len(strengths) < 3 or len(weaknesses) < 3:
        raise ValueError("강점과 보완할 점이 각각 3개 미만입니다.")
    if not model_answer:
        raise ValueError("모범 답변이 비었습니다.")
    return {
        "strengths": strengths[:3],
        "weaknesses": weaknesses[:3],
        "model_answer": model_answer,
    }


# ---------------------------------------------------------------------------
# 면접 단계별 API 작업
# ---------------------------------------------------------------------------


def context_block() -> str:
    ctx = st.session_state.context
    return prompts.build_context_block(ctx["role"], ctx["company"], ctx["materials"])


def asked_questions() -> list[str]:
    questions = []
    for record in st.session_state.records:
        questions.append(record["question"])
        if record["followup"]:
            questions.append(record["followup"]["question"])
    return questions


def generate_question(client: LLMClient) -> str:
    index = len(st.session_state.records) + 1
    total = st.session_state.context["num_questions"]
    user_prompt = prompts.build_question_prompt(context_block(), asked_questions(), index, total)
    return clean_question(call_llm_text(client, prompts.QUESTION_SYSTEM, user_prompt))


def generate_followup(client: LLMClient, question: str, answer: str) -> str:
    user_prompt = prompts.build_followup_prompt(context_block(), question, answer)
    return clean_question(call_llm_text(client, prompts.FOLLOWUP_SYSTEM, user_prompt))


def evaluate_answer(
    client: LLMClient, question: str, answer: str, is_followup: bool
) -> dict[str, Any]:
    user_prompt = prompts.build_evaluation_prompt(context_block(), question, answer, is_followup)
    return call_llm_json(client, prompts.EVALUATION_SYSTEM, user_prompt, validate_evaluation)


# ---------------------------------------------------------------------------
# 상태 관리
# ---------------------------------------------------------------------------


def init_state() -> None:
    defaults: dict[str, Any] = {
        "stage": "setup",  # setup → interview → report
        "context": None,
        "messages": [],  # 채팅 화면에 그릴 메시지 {"role", "kind", ...}
        "records": [],  # 본 질문별 기록
        "current": None,  # {"record_index", "kind": "main" | "followup"}
        "pending_answer": None,  # API 실패로 처리하지 못한 답변
        "report": None,
    }
    for key, value in defaults.items():
        st.session_state.setdefault(key, value)


def reset_interview(context: dict[str, Any]) -> None:
    st.session_state.update(
        stage="interview",
        context=context,
        messages=[],
        records=[],
        current=None,
        pending_answer=None,
        report=None,
    )


def back_to_setup() -> None:
    """설정 화면으로 돌아간다. 입력했던 context는 다음 설정의 기본값으로 남긴다."""
    st.session_state.update(
        stage="setup",
        messages=[],
        records=[],
        current=None,
        pending_answer=None,
        report=None,
    )


def add_message(role: str, kind: str, **payload: Any) -> None:
    """kind: question | answer | evaluation | reanswer | notice"""
    st.session_state.messages.append({"role": role, "kind": kind, **payload})


def new_record(question: str) -> dict[str, Any]:
    return {"question": question, "attempts": [], "followup": None}


def ask_main_question(question: str) -> None:
    st.session_state.records.append(new_record(question))
    index = len(st.session_state.records)
    st.session_state.current = {"record_index": index - 1, "kind": "main"}
    total = st.session_state.context["num_questions"]
    add_message("assistant", "question", label=f"질문 {index} / {total}", text=question)


def finish_interview() -> None:
    st.session_state.current = None
    st.session_state.stage = "report"
    add_message("assistant", "notice", text="마지막 질문까지 끝났어요.")


def move_to_next(next_question: str | None) -> None:
    if next_question is None:
        finish_interview()
    else:
        ask_main_question(next_question)


def has_more_questions() -> bool:
    return len(st.session_state.records) < st.session_state.context["num_questions"]


def process_answer(client: LLMClient, answer: str) -> None:
    """답변을 평가하고 다음 단계로 넘어간다.

    필요한 API 호출을 모두 끝낸 뒤 상태를 바꾼다. 중간에 호출이 실패하면
    상태가 그대로 남아 같은 답변으로 다시 시도할 수 있다.
    """
    current = st.session_state.current
    record = st.session_state.records[current["record_index"]]

    if current["kind"] == "followup":
        followup = record["followup"]
        evaluation = evaluate_answer(client, followup["question"], answer, is_followup=True)
        next_question = generate_question(client) if has_more_questions() else None

        add_message("user", "answer", text=answer)
        followup.update(answer=answer, evaluation=evaluation)
        add_message("assistant", "evaluation", label="꼬리 질문 평가", evaluation=evaluation)
        move_to_next(next_question)
        return

    evaluation = evaluate_answer(client, record["question"], answer, is_followup=False)
    reanswers_used = len(record["attempts"])
    passed = evaluation["average"] >= PASS_SCORE
    can_reanswer = not passed and reanswers_used < MAX_REANSWERS
    followup_question = None
    next_question = None
    if passed:
        followup_question = generate_followup(client, record["question"], answer)
    elif not can_reanswer and has_more_questions():
        next_question = generate_question(client)

    add_message("user", "answer", text=answer)
    record["attempts"].append({"answer": answer, "evaluation": evaluation})
    add_message("assistant", "evaluation", label="답변 평가", evaluation=evaluation)

    if passed:
        record["followup"] = {
            "question": followup_question,
            "answer": None,
            "evaluation": None,
        }
        st.session_state.current = {
            "record_index": current["record_index"],
            "kind": "followup",
        }
        add_message("assistant", "question", label="꼬리 질문", text=followup_question)
    elif can_reanswer:
        add_message(
            "assistant",
            "reanswer",
            evaluation=evaluation,
            question=record["question"],
            attempt=reanswers_used + 1,
        )
    else:
        add_message(
            "assistant", "notice", text="재답변 기회를 모두 썼어요. 다음 질문으로 넘어갈게요."
        )
        move_to_next(next_question)


# ---------------------------------------------------------------------------
# 표시용 포맷
# ---------------------------------------------------------------------------


def md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def format_score(score: float) -> str:
    return f"{score:g}"


# ---------------------------------------------------------------------------
# 최종 리포트
# ---------------------------------------------------------------------------


def scored_answers() -> list[dict[str, Any]]:
    """리포트에 쓰는 답변 목록. 본 질문은 마지막 답변, 꼬리 질문은 답한 경우만 넣는다."""
    rows = []
    for index, record in enumerate(st.session_state.records, start=1):
        if record["attempts"]:
            last = record["attempts"][-1]
            rows.append(
                {
                    "label": f"Q{index}",
                    "question": record["question"],
                    "answer": last["answer"],
                    "evaluation": last["evaluation"],
                    "attempts": len(record["attempts"]),
                }
            )
        followup = record["followup"]
        if followup and followup["evaluation"]:
            rows.append(
                {
                    "label": f"Q{index}-꼬리",
                    "question": followup["question"],
                    "answer": followup["answer"],
                    "evaluation": followup["evaluation"],
                    "attempts": 1,
                }
            )
    return rows


def evaluator_averages(rows: list[dict[str, Any]]) -> dict[str, float]:
    averages = {}
    for evaluator in prompts.EVALUATORS:
        name = evaluator["name"]
        scores = [
            e["score"]
            for row in rows
            for e in row["evaluation"]["evaluations"]
            if e["evaluator"] == name
        ]
        averages[name] = round(sum(scores) / len(scores), 1) if scores else 0.0
    return averages


def build_transcript(rows: list[dict[str, Any]]) -> str:
    blocks = []
    for row in rows:
        scores = "\n".join(
            f"  - {e['evaluator']} {format_score(e['score'])}점: {e['reason']}"
            for e in row["evaluation"]["evaluations"]
        )
        blocks.append(
            f"[{row['label']}] 질문: {row['question']}\n"
            f"답변: {row['answer']}\n"
            f"평균: {format_score(row['evaluation']['average'])}\n{scores}"
        )
    return "\n\n".join(blocks)


def generate_report(client: LLMClient) -> dict[str, Any]:
    rows = scored_answers()
    target = min(rows, key=lambda row: row["evaluation"]["average"])
    user_prompt = prompts.build_report_prompt(
        context_block(), build_transcript(rows), target["question"], target["answer"]
    )
    feedback = call_llm_json(client, prompts.REPORT_SYSTEM, user_prompt, validate_report)
    return {**feedback, "target": target}


def build_report_markdown(report: dict[str, Any]) -> str:
    rows = scored_answers()
    names = [e["name"] for e in prompts.EVALUATORS]
    ctx = st.session_state.context

    lines = [
        "# AI 면접 코치 최종 리포트",
        "",
        f"- 직군: {ctx['role']}",
        f"- 질문 수: {ctx['num_questions']}개",
        "",
        "## 질문별 점수표",
        "",
        "본 질문은 마지막 답변 기준, 꼬리 질문은 답한 경우만 표시합니다.",
        "",
        f"| 번호 | 질문 | 답변 횟수 | {' | '.join(names)} | 평균 |",
        f"|---|---|---|{'---|' * len(names)}---|",
    ]
    for row in rows:
        scores = {
            e["evaluator"]: format_score(e["score"]) for e in row["evaluation"]["evaluations"]
        }
        lines.append(
            f"| {row['label']} | {md_cell(row['question'])} | {row['attempts']} "
            f"| {' | '.join(scores[name] for name in names)} "
            f"| {format_score(row['evaluation']['average'])} |"
        )

    averages = evaluator_averages(rows)
    overall = round(sum(averages.values()) / len(averages), 1)
    lines += [
        "",
        "## 평가위원별 평균",
        "",
        "| 평가위원 | 평균 |",
        "|---|---|",
        *[f"| {name} | {format_score(score)} |" for name, score in averages.items()],
        f"| **전체** | **{format_score(overall)}** |",
        "",
        "## 강점",
        "",
        *[f"{i}. {s}" for i, s in enumerate(report["strengths"], start=1)],
        "",
        "## 보완할 점",
        "",
        *[f"{i}. {w}" for i, w in enumerate(report["weaknesses"], start=1)],
        "",
        "## 모범 답변 예시",
        "",
        f"가장 낮은 점수를 받은 {report['target']['label']} 답변을 다시 썼습니다.",
        "",
        f"**질문**: {report['target']['question']}",
        "",
        "**원래 답변**",
        "",
        *[f"> {line}" for line in report["target"]["answer"].splitlines()],
        "",
        "**모범 답변**",
        "",
        report["model_answer"],
        "",
        "## 전체 문답 기록",
        "",
    ]
    for row in rows:
        lines += [f"### {row['label']}. {row['question']}", "", row["answer"], ""]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 화면
# ---------------------------------------------------------------------------
# 표시 전용 요소는 HTML 문자열로 만들어 st.html로 그린다. 사용자·모델 텍스트는
# 반드시 html_text()로 이스케이프한다. 입력은 Streamlit 위젯을 그대로 쓴다.

STYLE_PATH = Path(__file__).with_name("style.css")
VOICE_SCRIPT_PATH = Path(__file__).with_name("voice_input.js")
STEPS = [("setup", "정보 입력"), ("interview", "면접"), ("report", "결과")]


def html_text(text: str) -> str:
    return escape(text).replace("\n", "<br>")


def short_name(evaluator: str) -> str:
    return evaluator.split()[0]


def score_bar(score: float) -> str:
    width = max(0.0, min(score, 10.0)) * 10
    level = "" if score >= PASS_SCORE else " low"
    return f'<span class="bar{level}"><span style="width:{width:.0f}%"></span></span>'


def status_pill(average: float) -> str:
    if average >= PASS_SCORE:
        return '<span class="pill pill-pass">기준 통과</span>'
    return '<span class="pill pill-fail">기준 미달</span>'


def inject_assets() -> None:
    """스타일과 음성 입력 스크립트를 넣는다.

    스크립트 실행은 저장소에 있는 고정 파일에만 허용한다. 사용자·모델 텍스트는 넣지 않는다.
    """
    st.html(f"<style>{STYLE_PATH.read_text(encoding='utf-8')}</style>")
    st.html(
        f"<script>{VOICE_SCRIPT_PATH.read_text(encoding='utf-8')}</script>",
        unsafe_allow_javascript=True,
    )


def interview_status() -> str:
    current = st.session_state.current
    record = st.session_state.records[current["record_index"]]
    if current["kind"] == "followup":
        return "꼬리 질문"
    if record["attempts"]:
        return f"재답변 {len(record['attempts'])} / {MAX_REANSWERS}"
    return "본 질문"


def answer_placeholder() -> str:
    status = interview_status()
    if status == "본 질문":
        return "답변을 입력하거나 마이크를 눌러 말해 주세요"
    if status == "꼬리 질문":
        return "꼬리 질문에 답해 주세요"
    return f"같은 질문에 다시 답해 주세요 ({status})"


def show_api_error(error: LLMError) -> None:
    message = str(error)
    if isinstance(error, llm_client.LLMAuthError):
        key_name = llm_client.API_KEY_NAMES.get(load_provider(), "API 키")
        message += f" .streamlit/secrets.toml 또는 .env의 {key_name}를 확인해 주세요."
    st.error(message)


GEMINI_NOTICE_TITLE = "실제 개인 정보는 넣지 마세요"
GEMINI_NOTICE_BODY = (
    "지금은 Gemini API로 질문하고 채점해요. Gemini 무료 등급에서는 입력한 내용을 Google이 "
    "서비스 개선에 쓰고, 사람이 직접 읽을 수도 있어요. 실제 이력서나 연락처 대신 "
    "test_data 폴더의 가상 자료로 테스트해 주세요."
)


def provider_notice_html(provider: str) -> str:
    if provider != "gemini":
        return ""
    return (
        '<div class="notice" role="note">'
        f'<p class="notice-title">{GEMINI_NOTICE_TITLE}</p>'
        f'<p class="notice-body">{GEMINI_NOTICE_BODY}</p>'
        "</div>"
    )


# --- 상단 내비게이션 ---


def render_topnav() -> None:
    stage = st.session_state.stage
    steps = "".join(
        f'<li class="{"on" if key == stage else ""}"><span>{number}</span>{label}</li>'
        for number, (key, label) in enumerate(STEPS, start=1)
    )
    with st.container(key="topnav", horizontal=True, vertical_alignment="center"):
        st.html('<p class="brand"><b>AI</b> 면접 코치</p>', width="content")
        with st.container(
            horizontal=True, vertical_alignment="center", width="content", gap="medium"
        ):
            st.html(f'<ol class="steps" aria-label="진행 단계">{steps}</ol>', width="content")
            if stage != "setup":
                render_restart()


def render_restart() -> None:
    with st.popover("새 면접", type="tertiary"):
        if st.session_state.stage == "interview":
            st.markdown("지금까지 답한 내용은 사라져요.")
        else:
            st.markdown("이 결과는 다시 볼 수 없어요. 필요하면 먼저 내려받아 주세요.")
        if st.button("새로 시작하기", type="primary", key="restart"):
            back_to_setup()
            st.rerun()


# --- 1. 정보 입력 ---


def render_setup(client: LLMClient, provider: str) -> None:
    with st.container(key="hero"):
        st.html(
            '<p class="hero-eyebrow">AI 모의 면접</p>'
            '<h1 class="hero-title">내 이력서로 준비하는<br><em>실전 모의 면접</em></h1>'
            '<p class="hero-sub">지원 직군과 회사 정보, 내 자료를 넣으면 질문을 하나씩 드려요.<br>'
            "답할 때마다 평가위원 4명이 점수와 고칠 점을 알려 드려요.</p>"
        )
    notice = provider_notice_html(provider)
    if notice:
        st.html(notice)
    form_col, side_col = st.columns([3, 2], gap="large")
    with form_col:
        render_setup_form(client)
    with side_col:
        st.html(setup_side_html())


def render_setup_form(client: LLMClient) -> None:
    previous = st.session_state.context or {}
    previous_role = previous.get("role", "")
    if previous_role in ROLE_OPTIONS:
        default_choice = previous_role
    else:
        default_choice = CUSTOM_ROLE if previous_role else ROLE_OPTIONS[0]

    st.html('<h2 class="section-title">면접 정보를 알려 주세요</h2>')
    role = st.segmented_control("지원 직군", ROLE_OPTIONS, default=default_choice, required=True)
    if role == CUSTOM_ROLE:
        role = st.text_input(
            "직군 이름",
            value=previous_role if default_choice == CUSTOM_ROLE else "",
            placeholder="예: 데이터 엔지니어",
        ).strip()
    company = st.text_area(
        "지원 회사 정보",
        value=previous.get("company", ""),
        height=180,
        placeholder="인재상, 사업 내용, 직무 설명을 붙여넣어 주세요",
    ).strip()
    files = st.file_uploader(
        "내 자료",
        type=SUPPORTED_TYPES,
        accept_multiple_files=True,
        help="자소서, 이력서, 포트폴리오를 올려 주세요. 스캔한 이미지 PDF는 글자를 읽을 수 없어요.",
    )
    previous_materials = previous.get("materials", [])
    if not files and previous_materials:
        st.caption(f"새로 올리지 않으면 지난번 자료 {len(previous_materials)}개를 그대로 써요.")
    num_questions = st.segmented_control(
        "질문 개수",
        [5, 10],
        default=previous.get("num_questions", 5),
        required=True,
        format_func=lambda n: f"{n}개",
    )

    if st.button("면접 시작하기", type="primary", width="stretch", key="start"):
        if not role:
            st.error("직군 이름을 입력해 주세요.")
            return
        if files:
            materials, warnings = read_uploaded_files(files)
        else:
            materials, warnings = previous_materials, []
        start_interview(
            client,
            {
                "role": role,
                "company": company,
                "materials": materials,
                "warnings": warnings,
                "num_questions": num_questions,
            },
        )


def setup_side_html() -> str:
    pass_score = format_score(PASS_SCORE)
    judges = "".join(
        f"<li><strong>{e['name']}</strong><span>{e['focus']}</span></li>"
        for e in prompts.EVALUATORS
    )
    return (
        '<div class="card card-blue">'
        '<p class="card-eyebrow">이렇게 진행돼요</p>'
        f'<p class="card-title">평균 {pass_score}점을 넘으면<br>꼬리 질문이 이어져요</p>'
        '<p class="card-body">'
        f"점수가 부족하면 같은 질문에 최대 {MAX_REANSWERS}번 다시 답할 수 있어요. "
        "마지막에는 점수표와 모범 답변을 정리해 드려요.</p>"
        "</div>"
        '<div class="card"><p class="card-head">평가위원 4명이 함께 봐요</p>'
        f'<ul class="judge-list">{judges}</ul></div>'
    )


def start_interview(client: LLMClient, context: dict[str, Any]) -> None:
    previous = {
        key: st.session_state[key]
        for key in ("stage", "context", "messages", "records", "current", "report")
    }
    reset_interview(context)
    try:
        with st.spinner("첫 질문을 준비하고 있어요"):
            question = generate_question(client)
    except LLMError as error:
        st.session_state.update(previous)
        show_api_error(error)
        return
    ask_main_question(question)
    st.rerun()


# --- 2. 면접 ---


def materials_summary(materials: list[dict[str, str]]) -> str:
    names = [m["name"] for m in materials]
    if not names:
        return "자료 없이 직군과 회사 정보로 질문하고 있어요"
    if len(names) == 1:
        return f"{names[0]}를 보고 질문하고 있어요"
    return f"{names[0]} 외 {len(names) - 1}개 자료를 보고 질문하고 있어요"


def interview_head_html(provider: str) -> str:
    ctx = st.session_state.context
    index = st.session_state.current["record_index"] + 1
    total = ctx["num_questions"]
    progress = index / total * 100
    return (
        '<div class="interview-head">'
        '<div class="ih-top">'
        f'<p class="ih-step"><strong>질문 {index}</strong>/ {total}</p>'
        f'<span class="ih-status">{interview_status()}</span>'
        "</div>"
        f'<div class="progress"><span style="width:{progress:.0f}%"></span></div>'
        f'<p class="ih-meta">{html_text(ctx["role"])} · '
        f"{html_text(materials_summary(ctx['materials']))}</p>"
        f"{interview_notice_html(provider)}"
        "</div>"
    )


def interview_notice_html(provider: str) -> str:
    if provider != "gemini":
        return ""
    return f'<p class="ih-notice">Gemini API로 채점 중이에요 · {GEMINI_NOTICE_TITLE}</p>'


def question_html(label: str, text: str) -> str:
    return (
        '<div class="msg-q">'
        f'<p class="msg-label">{html_text(label)}</p>'
        f'<p class="msg-q-text">{html_text(text)}</p>'
        "</div>"
    )


def evaluation_html(label: str, evaluation: dict[str, Any]) -> str:
    rows = "".join(
        "<li>"
        f'<div class="eval-row"><strong>{e["evaluator"]}</strong>'
        f'<span class="eval-score">{format_score(e["score"])}<small>/ 10</small></span></div>'
        f"{score_bar(e['score'])}"
        f'<p class="eval-reason">{html_text(e["reason"])}</p>'
        f'<p class="eval-tip"><span class="tag">개선</span>{html_text(e["improvement"])}</p>'
        "</li>"
        for e in evaluation["evaluations"]
    )
    return (
        '<div class="eval-card">'
        '<div class="eval-head"><div>'
        f'<p class="msg-label">{html_text(label)} · 평균</p>'
        f'<p class="eval-avg"><strong>{format_score(evaluation["average"])}</strong>/ 10</p>'
        f"</div>{status_pill(evaluation['average'])}</div>"
        f'<ul class="eval-list">{rows}</ul>'
        "</div>"
    )


def reanswer_html(message: dict[str, Any]) -> str:
    evaluation = message["evaluation"]
    body = (
        f"평균 {format_score(evaluation['average'])}점으로 "
        f"기준 {format_score(PASS_SCORE)}점에 못 미쳤어요."
    )
    if evaluation["gap"]:
        body += f" {evaluation['gap']}"
    return (
        '<div class="note">'
        '<p class="note-title">같은 질문에 다시 답해 주세요'
        f"<span>{message['attempt']} / {MAX_REANSWERS}</span></p>"
        f'<p class="note-body">{html_text(body)}</p>'
        f'<p class="note-q">{html_text(message["question"])}</p>'
        "</div>"
    )


def message_html(message: dict[str, Any]) -> str:
    kind = message["kind"]
    if kind == "question":
        return question_html(message["label"], message["text"])
    if kind == "answer":
        return f'<div class="msg-a">{html_text(message["text"])}</div>'
    if kind == "evaluation":
        return evaluation_html(message["label"], message["evaluation"])
    if kind == "reanswer":
        return reanswer_html(message)
    return f'<p class="msg-notice">{html_text(message["text"])}</p>'


def render_messages() -> None:
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.html(message_html(message))


def submit_answer(client: LLMClient, answer: str) -> None:
    try:
        with st.spinner("평가위원들이 답변을 보고 있어요"):
            process_answer(client, answer)
    except LLMError as error:
        st.session_state.pending_answer = answer
        show_api_error(error)
        return
    st.session_state.pending_answer = None
    st.rerun()


def render_interview(client: LLMClient, provider: str) -> None:
    column = st.container(key="narrow")
    with column:
        st.html(interview_head_html(provider))
        for warning in st.session_state.context["warnings"]:
            st.warning(warning, icon=":material/warning:")
        render_messages()
    # 새 답변을 먼저 처리해야 실패한 실행에서 바로 '다시 시도' 버튼이 보인다
    answer = st.chat_input(answer_placeholder())
    with column:
        if answer and answer.strip():
            submit_answer(client, answer.strip())
        pending = st.session_state.pending_answer
        if pending:
            st.warning("보낸 답변을 처리하지 못했어요.", icon=":material/error:")
            with st.expander("보낸 답변 보기"):
                st.text(pending)
            if st.button("같은 답변으로 다시 시도", icon=":material/refresh:"):
                submit_answer(client, pending)


# --- 3. 결과 ---


def judge_scores_html(averages: dict[str, float]) -> str:
    cards = "".join(
        '<div class="stat">'
        f'<p class="stat-label">{name}</p>'
        f'<p class="stat-value">{format_score(score)}<span>/ 10</span></p>'
        f"{score_bar(score)}"
        "</div>"
        for name, score in averages.items()
    )
    return (
        '<section class="section"><h2 class="section-title">평가위원별 점수</h2>'
        f'<div class="stat-grid">{cards}</div></section>'
    )


def score_list_html(rows: list[dict[str, Any]]) -> str:
    items = []
    for row in rows:
        evaluation = row["evaluation"]
        scores = " · ".join(
            f"{short_name(e['evaluator'])} {format_score(e['score'])}"
            for e in evaluation["evaluations"]
        )
        result = "pass" if evaluation["average"] >= PASS_SCORE else "fail"
        items.append(
            "<li><div>"
            f'<p class="q-title"><span class="chip">{row["label"]}</span>'
            f"{html_text(row['question'])}</p>"
            f'<p class="q-meta">답변 {row["attempts"]}회 · {scores}</p>'
            f'</div><p class="q-score {result}">'
            f"{format_score(evaluation['average'])}<span>점</span></p></li>"
        )
    return (
        '<section class="section"><h2 class="section-title">질문별 점수</h2>'
        '<p class="section-sub">본 질문은 마지막 답변 기준이에요. '
        "꼬리 질문은 답한 경우만 넣었어요.</p>"
        f'<ul class="score-list">{"".join(items)}</ul></section>'
    )


def feedback_html(report: dict[str, Any]) -> str:
    strengths = "".join(f"<li>{html_text(s)}</li>" for s in report["strengths"])
    weaknesses = "".join(f"<li>{html_text(w)}</li>" for w in report["weaknesses"])
    return (
        '<section class="section"><div class="two-col">'
        '<div class="card"><p class="card-head">강점</p>'
        f'<ol class="num-list">{strengths}</ol></div>'
        '<div class="card"><p class="card-head">보완할 점</p>'
        f'<ol class="num-list warn">{weaknesses}</ol></div>'
        "</div></section>"
    )


def model_answer_html(report: dict[str, Any]) -> str:
    target = report["target"]
    return (
        '<section class="section"><h2 class="section-title">모범 답변 예시</h2>'
        f'<p class="section-sub">점수가 가장 낮았던 {target["label"]} 답변'
        f"(평균 {format_score(target['evaluation']['average'])}점)을 다시 써 봤어요. "
        "대괄호 부분은 내 사례로 채워 주세요.</p>"
        f'<p class="ma-q">{html_text(target["question"])}</p>'
        '<div class="two-col">'
        '<div class="compare-box before"><p class="compare-label">원래 답변</p>'
        f'<p class="compare-text">{html_text(target["answer"])}</p></div>'
        '<div class="compare-box after"><p class="compare-label">다시 쓴 답변</p>'
        f'<p class="compare-text">{html_text(report["model_answer"])}</p></div>'
        "</div></section>"
    )


def render_report(client: LLMClient) -> None:
    rows = scored_answers()
    if not rows:
        st.info("채점한 답변이 없어 결과를 만들 수 없어요.")
        return

    if st.session_state.report is None:
        try:
            with st.spinner("결과를 정리하고 있어요"):
                st.session_state.report = generate_report(client)
        except LLMError as error:
            show_api_error(error)
            st.button("결과 다시 만들기", icon=":material/refresh:")
            return

    report = st.session_state.report
    ctx = st.session_state.context
    averages = evaluator_averages(rows)
    overall = round(sum(averages.values()) / len(averages), 1)

    with st.container(key="hero"):
        st.html(
            '<p class="hero-eyebrow">면접 결과</p>'
            f'<h1 class="hero-title">평균 <em>{format_score(overall)}점</em>이에요</h1>'
            f'<p class="hero-sub">{html_text(ctx["role"])} · 본 질문 {ctx["num_questions"]}개 · '
            f"답변 {len(rows)}개를 채점했어요</p>"
        )
        st.download_button(
            "결과 내려받기 (.md)",
            data=build_report_markdown(report),
            file_name="interview_report.md",
            mime="text/markdown",
            icon=":material/download:",
        )
    st.html(judge_scores_html(averages))
    st.html(score_list_html(rows))
    st.html(feedback_html(report))
    st.html(model_answer_html(report))
    with st.expander("면접 대화 다시 보기"):
        render_messages()


def main() -> None:
    st.set_page_config(
        page_title="AI 면접 코치",
        page_icon=":material/record_voice_over:",
        layout="wide",
    )
    inject_assets()
    init_state()
    render_topnav()

    provider = load_provider()
    if provider not in llm_client.PROVIDERS:
        st.error(
            f"LLM_PROVIDER 값 '{provider}'를 알 수 없어요. "
            f"{' 또는 '.join(llm_client.PROVIDERS)}로 설정해 주세요."
        )
        st.stop()
    key_name = llm_client.API_KEY_NAMES[provider]
    api_key = read_setting(key_name)
    if not api_key:
        guide = (
            " Gemini 키는 Google AI Studio에서 무료로 받을 수 있어요."
            if provider == "gemini"
            else ""
        )
        st.error(
            f"{key_name}가 없어요. .streamlit/secrets.toml 또는 .env에 키를 넣고 "
            f"다시 실행해 주세요.{guide}"
        )
        st.stop()
    client = get_client(provider, api_key)

    stage = st.session_state.stage
    if stage == "setup":
        render_setup(client, provider)
    elif stage == "interview":
        render_interview(client, provider)
    else:
        render_report(client)


if __name__ == "__main__":
    main()
